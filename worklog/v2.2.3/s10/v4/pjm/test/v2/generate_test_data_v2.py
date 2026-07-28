from __future__ import annotations

import csv
import importlib.util
import json
import os
import re
import sys
import tempfile
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


OUTPUT_DIR = Path(__file__).resolve().parent
BASE_GENERATOR = OUTPUT_DIR.parent / "generate_test_data.py"
VERSION = "pjm-test-v2"
SNAPSHOT_THROUGH = date(2027, 6, 30)
WORKBOOK_TIMESTAMP = datetime(2027, 6, 30, 0, 0, 0)
WORKBOOK_TIMESTAMP_XML = b"2027-06-30T00:00:00Z"

TABLE_KEYS = {
    "ods_project_subject_domain_v2": ("project_no", "node_task"),
    "ods_progress_measure_v2": ("project_no", "node_task"),
    "ods_quality_issue_v2": ("project_no", "issue_name"),
    "ods_quality_measure_v2": ("project_no", "issue_name"),
    "ods_tech_state_v2": ("project_no", "tech_state_name"),
    "ods_tech_state_measure_v2": ("project_no", "tech_state_name"),
    "ods_risk_info_v2": ("project_no", "risk_name"),
    "ods_risk_measure_v2": ("project_no", "risk_name"),
    "ods_material_info_v2": ("project_no", "pbs_no"),
    "ods_budget_v2": ("budget_no",),
}


def load_base_generator() -> Any:
    spec = importlib.util.spec_from_file_location("pjm_test_data_base", BASE_GENERATOR)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"无法加载基础生成脚本: {BASE_GENERATOR}")
    module = importlib.util.module_from_spec(spec)
    previous = sys.dont_write_bytecode
    sys.dont_write_bytecode = True
    try:
        spec.loader.exec_module(module)
    finally:
        sys.dont_write_bytecode = previous
    return module


BASE = load_base_generator()


def parse_date(value: str) -> date:
    if not value:
        raise ValueError("日期不能为空")
    return date.fromisoformat(value)


def set_date_pair(row: dict[str, str], date_field: str, week_field: str, value: date | None) -> None:
    row[date_field] = BASE.fmt(value)
    row[week_field] = BASE.week_of(value)


def find_row(
    rows: list[dict[str, str]],
    project_no: str,
    field: str,
    suffix: str,
) -> dict[str, str]:
    matches = [
        row
        for row in rows
        if row.get("project_no") == project_no and row.get(field, "").endswith(suffix)
    ]
    if len(matches) != 1:
        raise AssertionError(
            f"期望唯一记录: project_no={project_no}, {field}=*{suffix}; 实际 {len(matches)} 条"
        )
    return matches[0]


def close_progress_node(
    rows: list[dict[str, str]],
    project_no: str,
    task_suffix: str,
    actual: date,
    changed_plan: bool,
) -> None:
    row = find_row(rows, project_no, "node_task", task_suffix)
    plan = parse_date(row["plan_date"])
    delay_days = max((actual - plan).days, 0)
    row["completion_status"] = "超期已完成已变更" if changed_plan else "超期已完成未变更"
    row["actual_start_date"] = row["actual_start_date"] or row["plan_start_date"]
    set_date_pair(row, "actual_date", "actual_week", actual)
    set_date_pair(row, "last_update_time", "last_update_week", actual)
    row["delay_expected_date"] = ""
    row["original_plan_date"] = row["original_plan_date"] or (row["plan_date"] if changed_plan else "")
    row["delay_days_changed"] = str(delay_days) if changed_plan else ""
    row["delay_days_unchanged"] = "" if changed_plan else str(delay_days)
    row["delay_applied"] = "是"
    row["incomplete_reason"] = "第二版快照：延期事项已完成闭环"
    row["risk_level"] = "中"
    row["risk_content"] = f"{task_suffix}已完成，后续窗口持续跟踪"
    row["delay_impact"] = f"后续节点窗口压缩{delay_days}天"
    row["highlight"] = "第二版快照已完成"


def complete_progress_on_time(
    rows: list[dict[str, str]],
    project_no: str,
    task_suffix: str,
    actual: date,
) -> None:
    row = find_row(rows, project_no, "node_task", task_suffix)
    row["completion_status"] = "按时完成"
    row["actual_start_date"] = row["actual_start_date"] or row["plan_start_date"]
    set_date_pair(row, "actual_date", "actual_week", actual)
    set_date_pair(row, "last_update_time", "last_update_week", actual)
    for field in (
        "delay_expected_date",
        "incomplete_reason",
        "risk_level",
        "risk_content",
        "delay_impact",
        "original_plan_date",
        "delay_days_changed",
        "delay_days_unchanged",
    ):
        row[field] = ""
    row["delay_applied"] = "否"
    row["highlight"] = "第二版快照：关键件按期交付"


def evolve_progress(rows: list[dict[str, str]]) -> None:
    close_progress_node(rows, "XM-2026-C01", "联调准备", date(2026, 7, 24), False)
    close_progress_node(rows, "XM-2026-C02", "联调准备", date(2026, 8, 28), True)
    complete_progress_on_time(rows, "XM-2026-C03", "关键件交付", date(2026, 7, 15))
    close_progress_node(rows, "XM-2026-C08", "关键件交付", date(2026, 12, 29), True)


def close_quality_issue(row: dict[str, str], completed: date) -> None:
    row["status"] = "已完成技术和管理归零"
    row["zero_plan"] = row["zero_plan"] or f"完成{row['issue_name']}归零并输出验证记录"
    row["zero_plan_synced"] = "是"
    row["new_plan_count"] = str(int(row["new_plan_count"] or "0") + 1)
    row["current_progress"] = "第二版快照：技术与管理归零均已完成"
    set_date_pair(row, "zero_complete_date", "zero_complete_week", completed)
    set_date_pair(row, "last_update_time", "last_update_week", completed)


def evolve_quality(rows: list[dict[str, str]]) -> None:
    completion_offsets = {
        "XM-2026-C02": 12,
        "XM-2026-C04": 14,
        "XM-2026-C08": 21,
    }
    for project_no, offset in completion_offsets.items():
        row = find_row(rows, project_no, "issue_name", "异常问题")
        baseline = parse_date(row["zero_complete_date"] or row["last_update_time"])
        close_quality_issue(row, baseline + timedelta(days=offset))


def close_tech_state(row: dict[str, str], completed: date) -> None:
    category = row["change_category"]
    signature_date = completed - timedelta(days=5)
    row["completion_signature"] = "是"
    set_date_pair(row, "signature_closure_date", "signature_closure_week", signature_date)
    row["review_situation"] = (
        "已提出需求，已签署"
        if category == "III"
        else "已评估评审已通过"
    )
    row["file_signature_status"] = (
        "已提出需求，已签署"
        if category == "III"
        else "已评估评审，已签署"
    )
    set_date_pair(row, "file_signature_date", "file_signature_week", signature_date)
    row["reform_status"] = "不涉及" if category == "III" else "已落实整改"
    set_date_pair(row, "reform_date", "reform_week", completed)
    row["plan_synced"] = "是"
    set_date_pair(row, "last_update_time", "last_update_week", completed)
    row["remark"] = "第二版快照：文件签署及整改已闭环"


def evolve_tech(rows: list[dict[str, str]]) -> None:
    for project_no, extra_days in {
        "XM-2026-C02": 5,
        "XM-2026-C04": 7,
        "XM-2026-C06": 3,
        "XM-2026-C10": 9,
    }.items():
        row = find_row(rows, project_no, "tech_state_name", row_suffix_for_project(project_no))
        completed = parse_date(row["plan_reform_date"]) + timedelta(days=extra_days)
        close_tech_state(row, completed)

    for project_no in ("XM-2026-C03", "XM-2026-C09"):
        row = find_row(rows, project_no, "tech_state_name", row_suffix_for_project(project_no))
        closure_date = parse_date(row["file_signature_date"])
        set_date_pair(row, "reform_date", "reform_week", closure_date)
        set_date_pair(row, "last_update_time", "last_update_week", closure_date)
        row["remark"] = "第二版快照：不涉及整改，签署日作为闭环日期"


def row_suffix_for_project(project_no: str) -> str:
    project = next(project for project in BASE.PROJECTS if project["project_no"] == project_no)
    return project["key_item"]


def release_risk(row: dict[str, str], released: date) -> None:
    row["risk_status"] = "已释放"
    set_date_pair(row, "risk_release_date", "risk_release_week", released)
    set_date_pair(row, "final_release_time", "final_release_week", released)
    set_date_pair(row, "progress_stat_time", "progress_stat_week", released)
    set_date_pair(row, "last_update_time", "last_update_week", released)
    row["release_plan_synced"] = "是"
    row["progress_situation"] = "第二版快照：措施已闭环，风险完成释放"
    row["remark"] = "第二版快照：释放证据已归档"


def evolve_risks(rows: list[dict[str, str]]) -> None:
    for project_no, days_after_submit in {
        "XM-2026-C02": 49,
        "XM-2026-C04": 56,
        "XM-2026-C08": 63,
    }.items():
        row = find_row(rows, project_no, "risk_name", "风险")
        release_risk(row, parse_date(row["risk_submit_time"]) + timedelta(days=days_after_submit))

    c06 = find_row(rows, "XM-2026-C06", "risk_name", "风险")
    c06["risk_level"] = "中"
    c06["progress_situation"] = "第二版快照：物料已到货，装配验证中，暂未达到释放条件"
    update = parse_date(c06["last_update_time"]) + timedelta(days=14)
    set_date_pair(c06, "progress_stat_time", "progress_stat_week", update)
    set_date_pair(c06, "last_update_time", "last_update_week", update)


def deliver_material(row: dict[str, str], delay_days: int) -> None:
    actual = parse_date(row["contract_delivery_date"]) + timedelta(days=delay_days)
    inspection = actual + timedelta(days=4)
    install = inspection + timedelta(days=7)
    set_date_pair(row, "actual_delivery_date", "actual_delivery_week", actual)
    set_date_pair(row, "complete_inspect_date", "complete_inspect_week", inspection)
    set_date_pair(row, "install_date", "install_week", install)
    set_date_pair(row, "last_update_time", "last_update_week", install)
    row["weekly_progress"] = "第二版快照：已完成到货检验并进入装配"
    row["affects_major_node"] = "否"
    row["risk_level"] = "低"
    row["risk_content"] = ""
    row["delay_impact"] = ""
    row["remark"] = "第二版快照：到货事项已闭环"


def evolve_materials(rows: list[dict[str, str]]) -> None:
    for project_no, delay_days in {
        "XM-2026-C03": 10,
        "XM-2026-C06": 18,
        "XM-2026-C08": 21,
    }.items():
        row = find_row(rows, project_no, "pbs_no", row_suffix_for_pbs(project_no))
        deliver_material(row, delay_days)


def row_suffix_for_pbs(project_no: str) -> str:
    project = next(project for project in BASE.PROJECTS if project["project_no"] == project_no)
    return project["pbs_no"]


def money(value: Decimal) -> str:
    return str(value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def evolve_budgets(rows: list[dict[str, str]]) -> None:
    for idx, row in enumerate(rows):
        budget = Decimal(row["budget_amount_adjusted"])
        delta = (budget * Decimal("0.04")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        target = ("prepaid_amount", "book_cost_amount", "payable_amount")[idx % 3]
        row[target] = money(Decimal(row[target]) + delta)

    rows.append(
        {
            "project_no": "XM-2026-C01",
            "budget_no": "YS-2026-C01-03",
            "subtopic": "综合统筹平台-补充联试验证",
            "research_lab": "总体室",
            "budget_amount_adjusted": "90.00",
            "prepaid_amount": "9.00",
            "book_cost_amount": "22.50",
            "payable_amount": "13.50",
        }
    )


def build_tables() -> tuple[dict[str, list[dict[str, str]]], dict[str, list[dict[str, str]]]]:
    base_progress = BASE.build_progress_rows()
    base_quality = BASE.build_quality_rows()
    base_tech = BASE.build_tech_rows()
    base_risk = BASE.build_risk_rows()
    base_material = BASE.build_material_rows()
    base_budget = BASE.build_budget_rows()
    base_tables = {
        "ods_project_subject_domain_v2": base_progress,
        "ods_progress_measure_v2": BASE.build_progress_measure_rows(base_progress),
        "ods_quality_issue_v2": base_quality,
        "ods_quality_measure_v2": BASE.build_quality_measure_rows(base_quality),
        "ods_tech_state_v2": base_tech,
        "ods_tech_state_measure_v2": BASE.build_tech_measure_rows(base_tech),
        "ods_risk_info_v2": base_risk,
        "ods_risk_measure_v2": BASE.build_risk_measure_rows(base_risk),
        "ods_material_info_v2": base_material,
        "ods_budget_v2": base_budget,
    }

    progress = BASE.build_progress_rows()
    quality = BASE.build_quality_rows()
    tech = BASE.build_tech_rows()
    risk = BASE.build_risk_rows()
    material = BASE.build_material_rows()
    budget = BASE.build_budget_rows()
    evolve_progress(progress)
    evolve_quality(quality)
    evolve_tech(tech)
    evolve_risks(risk)
    evolve_materials(material)
    evolve_budgets(budget)

    updated_tables = {
        "ods_project_subject_domain_v2": progress,
        "ods_progress_measure_v2": BASE.build_progress_measure_rows(progress),
        "ods_quality_issue_v2": quality,
        "ods_quality_measure_v2": BASE.build_quality_measure_rows(quality),
        "ods_tech_state_v2": tech,
        "ods_tech_state_measure_v2": BASE.build_tech_measure_rows(tech),
        "ods_risk_info_v2": risk,
        "ods_risk_measure_v2": BASE.build_risk_measure_rows(risk),
        "ods_material_info_v2": material,
        "ods_budget_v2": budget,
    }
    return normalize_tables(base_tables), normalize_tables(updated_tables)


def normalize_tables(
    tables: dict[str, list[dict[str, str]]],
) -> dict[str, list[dict[str, str]]]:
    return {
        table_name: [
            {
                field: row.get(field, "")
                for field in BASE.CSV_SCHEMAS[table_name]
            }
            for row in rows
        ]
        for table_name, rows in tables.items()
    }


def row_key(row: dict[str, str], fields: tuple[str, ...]) -> tuple[str, ...]:
    return tuple(row.get(field, "") for field in fields)


def assert_unique(table_name: str, rows: list[dict[str, str]]) -> None:
    fields = TABLE_KEYS[table_name]
    keys = [row_key(row, fields) for row in rows]
    if any(not all(key) for key in keys):
        raise AssertionError(f"{table_name} 存在空业务键: {fields}")
    duplicates = sorted({key for key in keys if keys.count(key) > 1})
    if duplicates:
        raise AssertionError(f"{table_name} 存在重复业务键: {duplicates[:5]}")


def validate_tables(tables: dict[str, list[dict[str, str]]]) -> None:
    if set(tables) != set(BASE.CSV_SCHEMAS):
        raise AssertionError("输出表集合与基础版本不一致")
    valid_projects = {project["project_no"] for project in BASE.PROJECTS}
    for table_name, rows in tables.items():
        assert_unique(table_name, rows)
        expected_fields = set(BASE.CSV_SCHEMAS[table_name])
        if any(set(row) != expected_fields for row in rows):
            raise AssertionError(f"{table_name} 字段集合与基础版本不一致")
        projects = {row["project_no"] for row in rows}
        if projects != valid_projects:
            raise AssertionError(f"{table_name} 项目覆盖不完整: {sorted(projects)}")
        validate_date_week_pairs(table_name, rows)

    paired_tables = (
        ("ods_project_subject_domain_v2", "ods_progress_measure_v2"),
        ("ods_quality_issue_v2", "ods_quality_measure_v2"),
        ("ods_tech_state_v2", "ods_tech_state_measure_v2"),
        ("ods_risk_info_v2", "ods_risk_measure_v2"),
    )
    for fact_name, measure_name in paired_tables:
        key_fields = TABLE_KEYS[fact_name]
        fact_keys = {row_key(row, key_fields) for row in tables[fact_name]}
        measure_keys = {row_key(row, key_fields) for row in tables[measure_name]}
        if fact_name == "ods_project_subject_domain_v2":
            if not measure_keys.issubset(fact_keys):
                raise AssertionError("进度措施存在无法关联的项目节点")
        elif fact_keys != measure_keys:
            raise AssertionError(f"{measure_name} 与 {fact_name} 业务键不一致")

    for table_name in (
        "ods_progress_measure_v2",
        "ods_quality_measure_v2",
        "ods_tech_state_measure_v2",
        "ods_risk_measure_v2",
    ):
        for row in tables[table_name]:
            closure_fields = (
                "final_closure_date",
                "final_closure_week",
                "closure_deliverable_type",
                "closure_deliverable",
            )
            has_complete_evidence = all(bool(row[field]) for field in closure_fields)
            if (row["closure_status"] == "已闭环") != has_complete_evidence:
                raise AssertionError(
                    f"{table_name} 闭环状态与证据不一致: {row['project_no']}"
                )

    accepted_progress = {
        "按时完成",
        "正常待完成",
        "超期已完成未变更",
        "超期已完成已变更",
        "超期未完成未变更",
        "超期未完成已变更",
        "不正常待变更",
    }
    if {
        row["completion_status"] for row in tables["ods_project_subject_domain_v2"]
    } - accepted_progress:
        raise AssertionError("进度状态包含未支持枚举")
    completed_statuses = {"按时完成", "超期已完成未变更", "超期已完成已变更"}
    for row in tables["ods_project_subject_domain_v2"]:
        if (row["completion_status"] in completed_statuses) != bool(row["actual_date"]):
            raise AssertionError(
                f"进度完成状态与实际完成日期不一致: {row['project_no']} / {row['node_task']}"
            )

    for row in tables["ods_risk_info_v2"]:
        if (row["risk_status"] == "已释放") != bool(row["risk_release_date"]):
            raise AssertionError(
                f"风险释放状态与释放日期不一致: {row['project_no']} / {row['risk_name']}"
            )

    overrun_budget_nos: set[str] = set()
    for row in tables["ods_budget_v2"]:
        amounts = [
            Decimal(row[field])
            for field in (
                "budget_amount_adjusted",
                "prepaid_amount",
                "book_cost_amount",
                "payable_amount",
            )
        ]
        if any(value < 0 for value in amounts):
            raise AssertionError(f"预算金额不能为负数: {row['budget_no']}")
        if any(Decimal(row[field]).as_tuple().exponent != -2 for field in (
            "budget_amount_adjusted",
            "prepaid_amount",
            "book_cost_amount",
            "payable_amount",
        )):
            raise AssertionError(f"预算金额必须保留两位小数: {row['budget_no']}")
        executed = sum(amounts[1:])
        if executed > amounts[0]:
            overrun_budget_nos.add(row["budget_no"])

    if len(tables["ods_budget_v2"]) != 26:
        raise AssertionError("第二版预算应包含 26 条记录")
    expected_overruns = {
        "YS-2026-C02-01",
        "YS-2026-C04-02",
        "YS-2026-C06-03",
        "YS-2026-C09-01",
    }
    if overrun_budget_nos != expected_overruns:
        raise AssertionError(f"预算超支场景发生漂移: {sorted(overrun_budget_nos)}")


def validate_date_week_pairs(
    table_name: str,
    rows: list[dict[str, str]],
) -> None:
    observed_date_fields = {
        "actual_date",
        "last_update_time",
        "zero_complete_date",
        "signature_closure_date",
        "file_signature_date",
        "reform_date",
        "risk_release_date",
        "progress_stat_time",
        "actual_delivery_date",
        "complete_inspect_date",
        "install_date",
        "follow_up_date",
        "final_closure_date",
    }
    for row_index, row in enumerate(rows, start=2):
        for field, value in row.items():
            if not value:
                continue
            if field.endswith("_date"):
                week_field = f"{field[:-5]}_week"
            elif field.endswith("_time"):
                week_field = f"{field[:-5]}_week"
            else:
                continue
            if week_field not in row:
                continue
            parsed = parse_date(value)
            if row[week_field] != BASE.week_of(parsed):
                raise AssertionError(
                    f"{table_name} 第 {row_index} 行日期与周数不一致: {field}/{week_field}"
                )
            if field in observed_date_fields and parsed > SNAPSHOT_THROUGH:
                raise AssertionError(
                    f"{table_name} 第 {row_index} 行实际日期晚于快照截止日: {field}={value}"
                )


def change_stats(
    base_tables: dict[str, list[dict[str, str]]],
    updated_tables: dict[str, list[dict[str, str]]],
) -> dict[str, dict[str, int]]:
    result: dict[str, dict[str, int]] = {}
    for table_name, updated_rows in updated_tables.items():
        fields = TABLE_KEYS[table_name]
        base_by_key = {row_key(row, fields): row for row in base_tables[table_name]}
        updated_by_key = {row_key(row, fields): row for row in updated_rows}
        inserted = len(updated_by_key.keys() - base_by_key.keys())
        deleted = len(base_by_key.keys() - updated_by_key.keys())
        updated = sum(
            base_by_key[key] != updated_by_key[key]
            for key in base_by_key.keys() & updated_by_key.keys()
        )
        result[table_name] = {
            "rows": len(updated_rows),
            "inserted": inserted,
            "updated": updated,
            "deleted": deleted,
            "unchanged": len(updated_rows) - inserted - updated,
        }
    return result


def validate_row_order(
    base_tables: dict[str, list[dict[str, str]]],
    updated_tables: dict[str, list[dict[str, str]]],
) -> None:
    for table_name, base_rows in base_tables.items():
        fields = TABLE_KEYS[table_name]
        base_keys = [row_key(row, fields) for row in base_rows]
        updated_keys = [row_key(row, fields) for row in updated_tables[table_name]]
        if updated_keys[: len(base_keys)] != base_keys:
            raise AssertionError(f"{table_name} 原有业务行顺序发生变化")


def polish_workbook(output_dir: Path, table_name: str, row_count: int) -> None:
    path = output_dir / f"{table_name}.xlsx"
    workbook = load_workbook(path)
    try:
        workbook.properties.created = WORKBOOK_TIMESTAMP
        workbook.properties.modified = WORKBOOK_TIMESTAMP
        workbook.properties.lastModifiedBy = "pjm-test-v2-generator"
        sheet = workbook.active
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in sheet.iter_cols():
            values = [str(cell.value or "") for cell in column_cells[: min(row_count + 1, 101)]]
            width = min(max(max(map(len, values), default=8) + 2, 10), 36)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        workbook.save(path)
    finally:
        workbook.close()
    normalize_xlsx_archive(path)


def normalize_xlsx_archive(path: Path) -> None:
    normalized_path = path.with_name(f".{path.name}.normalized")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        normalized_path,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=9,
    ) as target:
        for source_info in sorted(source.infolist(), key=lambda item: item.filename):
            content = source.read(source_info.filename)
            if source_info.filename == "docProps/core.xml":
                for tag in (b"created", b"modified"):
                    pattern = (
                        rb"(<dcterms:"
                        + tag
                        + rb"\b[^>]*>)[^<]*(</dcterms:"
                        + tag
                        + rb">)"
                    )
                    content, replacement_count = re.subn(
                        pattern,
                        lambda match: match.group(1)
                        + WORKBOOK_TIMESTAMP_XML
                        + match.group(2),
                        content,
                    )
                    if replacement_count != 1:
                        raise AssertionError(
                            f"{path.name} 缺少可归一化的 {tag.decode()} 时间"
                        )
            normalized_info = zipfile.ZipInfo(
                filename=source_info.filename,
                date_time=(1980, 1, 1, 0, 0, 0),
            )
            normalized_info.compress_type = zipfile.ZIP_DEFLATED
            normalized_info.create_system = source_info.create_system
            normalized_info.external_attr = source_info.external_attr
            normalized_info.internal_attr = source_info.internal_attr
            normalized_info.flag_bits = source_info.flag_bits
            normalized_info.comment = source_info.comment
            target.writestr(
                normalized_info,
                content,
                compress_type=zipfile.ZIP_DEFLATED,
                compresslevel=9,
            )
    os.replace(normalized_path, path)


def verify_outputs(
    output_dir: Path,
    tables: dict[str, list[dict[str, str]]],
) -> None:
    for table_name, rows in tables.items():
        csv_path = output_dir / f"{table_name}.csv"
        with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            csv_rows = list(reader)
        if csv_rows[0] != BASE.CSV_SCHEMAS[table_name]:
            raise AssertionError(f"{csv_path.name} 表头不一致")
        if len(csv_rows) - 1 != len(rows):
            raise AssertionError(f"{csv_path.name} 行数不一致")

        workbook = load_workbook(output_dir / f"{table_name}.xlsx", read_only=True, data_only=True)
        try:
            sheet = workbook.active
            xlsx_headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
            if xlsx_headers != BASE.CSV_SCHEMAS[table_name]:
                raise AssertionError(f"{table_name}.xlsx 表头不一致")
            if sheet.max_row - 1 != len(rows):
                raise AssertionError(f"{table_name}.xlsx 行数不一致")
            xlsx_rows = [
                [str(cell.value or "") for cell in cells]
                for cells in sheet.iter_rows(min_row=2)
            ]
            if xlsx_rows != csv_rows[1:]:
                raise AssertionError(f"{table_name} 的 CSV 与 XLSX 数据内容不一致")
        finally:
            workbook.close()


def main() -> None:
    base_tables, tables = build_tables()
    validate_tables(tables)
    validate_row_order(base_tables, tables)

    stats = change_stats(base_tables, tables)
    manifest = {
        "version": VERSION,
        "snapshotThrough": SNAPSHOT_THROUGH.isoformat(),
        "source": "../generate_test_data.py",
        "purpose": "同一业务主键的第二次全量快照，同时覆盖更新与新增",
        "tables": stats,
        "businessChanges": [
            "4 个进度节点完成状态推进",
            "3 个质量问题完成技术和管理归零",
            "4 个技术状态完成签署或整改闭环，2 条既有闭环补齐日期",
            "3 个风险完成释放，1 个风险继续跟踪",
            "3 个重要物料完成到货检验与装配",
            "25 条预算三本账发生执行额变化，并新增 1 条预算",
        ],
    }

    with tempfile.TemporaryDirectory(prefix=".pjm-test-v2-", dir=OUTPUT_DIR) as temp_name:
        staging_dir = Path(temp_name)
        previous_test_dir = BASE.TEST_DIR
        try:
            BASE.TEST_DIR = staging_dir
            for table_name, rows in tables.items():
                BASE.write_table(table_name, rows)
                polish_workbook(staging_dir, table_name, len(rows))
            (staging_dir / "manifest.json").write_text(
                json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            verify_outputs(staging_dir, tables)
        finally:
            BASE.TEST_DIR = previous_test_dir

        artifact_names = [
            name
            for table_name in tables
            for name in (f"{table_name}.csv", f"{table_name}.xlsx")
        ]
        for name in artifact_names:
            os.replace(staging_dir / name, OUTPUT_DIR / name)
        os.replace(staging_dir / "manifest.json", OUTPUT_DIR / "manifest.json")

    verify_outputs(OUTPUT_DIR, tables)

    print(f"generated {VERSION}:")
    for table_name, table_stats in stats.items():
        print(
            f"{table_name}: rows={table_stats['rows']}, "
            f"updated={table_stats['updated']}, inserted={table_stats['inserted']}"
        )


if __name__ == "__main__":
    main()
